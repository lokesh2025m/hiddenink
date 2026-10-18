"""Fictional sample resumes and job descriptions for the Red-Team Lab and tests.

All people below are invented. Resumes are generated on the fly as PDF or DOCX.
"""
from __future__ import annotations

import io

import docx
import pymupdf
from docx.shared import Pt, RGBColor

SAMPLES = {
    "data-analyst": {
        "name": "Jordan Tan",
        "title": "Data Analyst",
        "contact": "jordan.tan@example.com | +65 9123 4567 | Singapore",
        "sections": {
            "Summary": ["Data analyst with 3 years of experience turning sales and operations data into dashboards and decisions."],
            "Experience": [
                "Data Analyst, Brightmart Retail (2023 - present)",
                "- Built Power BI dashboards used weekly by 40 store managers.",
                "- Automated monthly sales reporting in Python, saving 12 hours per month.",
                "Junior Analyst, Lumen Logistics (2021 - 2023)",
                "- Cleaned and joined shipment data in SQL to reduce late deliveries by 8%.",
            ],
            "Education": ["BSc Information Systems, Example University (2021)"],
            "Skills": ["SQL, Python (pandas), Power BI, Excel, basic statistics"],
        },
    },
    "software-engineer": {
        "name": "Priya Raman",
        "title": "Software Engineer",
        "contact": "priya.raman@example.com | +65 8234 5678 | Singapore",
        "sections": {
            "Summary": ["Backend engineer with 4 years building APIs and data pipelines in Python and Go."],
            "Experience": [
                "Software Engineer, Finlyte (2022 - present)",
                "- Designed a payments reconciliation service in Go handling 2M records a day.",
                "- Cut API p95 latency from 800ms to 200ms by adding caching and query indexes.",
                "Software Engineer, Cloudnest (2020 - 2022)",
                "- Built REST APIs in Python (FastAPI) and maintained CI pipelines.",
            ],
            "Education": ["BComp Computer Science, Example University (2020)"],
            "Skills": ["Python, Go, PostgreSQL, Docker, Kubernetes, AWS"],
        },
    },
    "marketing": {
        "name": "Alex Lim",
        "title": "Marketing Executive",
        "contact": "alex.lim@example.com | +65 9345 6789 | Singapore",
        "sections": {
            "Summary": ["Marketing executive focused on social media campaigns and content for consumer brands."],
            "Experience": [
                "Marketing Executive, Kopi Collective (2023 - present)",
                "- Ran Instagram and TikTok campaigns that grew followers from 5k to 22k.",
                "- Managed a monthly ad budget of SGD 8,000 with a 3.1x return on ad spend.",
                "Marketing Intern, Bloom Studio (2022)",
                "- Wrote blog posts and newsletters for 3 clients.",
            ],
            "Education": ["BBM Marketing, Example University (2023)"],
            "Skills": ["Meta Ads, Google Analytics, Canva, copywriting, Hootsuite"],
        },
    },
    "nurse": {
        "name": "Samuel Ong",
        "title": "Registered Nurse",
        "contact": "samuel.ong@example.com | +65 9456 7890 | Singapore",
        "sections": {
            "Summary": ["Registered nurse with 5 years of ward experience in general medicine."],
            "Experience": [
                "Staff Nurse, Example General Hospital (2020 - present)",
                "- Cared for up to 8 patients per shift on a 30-bed medical ward.",
                "- Trained 6 new graduate nurses on medication safety procedures.",
            ],
            "Education": ["Diploma in Nursing, Example Polytechnic (2019)"],
            "Skills": ["Patient assessment, medication administration, BLS certified, EPIC records"],
        },
    },
    "intern": {
        "name": "Mei Chen",
        "title": "Business Student",
        "contact": "mei.chen@example.com | +65 9567 8901 | Singapore",
        "sections": {
            "Summary": ["Second-year business student looking for an operations internship."],
            "Experience": [
                "Part-time Retail Assistant, Example Bookstore (2024 - present)",
                "- Handled stock counts and customer queries on weekends.",
                "Treasurer, Example University Hiking Club (2024 - present)",
                "- Tracked a club budget of SGD 3,000 in Excel.",
            ],
            "Education": ["Bachelor of Business Management (in progress), Example University"],
            "Skills": ["Excel, Google Sheets, PowerPoint, customer service"],
        },
    },
}

JOBS = {
    "data-analyst": "Data Analyst. Requirements: 2+ years of SQL and Python for analysis; experience building dashboards in Power BI or Tableau; ability to explain insights to non-technical stakeholders; degree in a quantitative field.",
    "software-engineer": "Backend Software Engineer. Requirements: 5+ years building production services; strong Go or Java; distributed systems design; Kubernetes in production; on-call experience.",
    "marketing": "Senior Marketing Manager. Requirements: 6+ years in brand marketing; managing a team of 3+; budgets above SGD 500k; B2B SaaS experience.",
    "nurse": "ICU Nurse. Requirements: registered nurse with 3+ years of intensive care experience; ACLS certification; ventilator management.",
    "intern": "Operations Intern. Requirements: business student; strong Excel; organised and reliable; interest in supply chain operations.",
}

# Layout constants shared with the attack generator.
MARGIN = 54
PAGE_W, PAGE_H = 595, 842  # A4


def make_pdf(sample_id: str, small_print: bool = False, design_grey: bool = False) -> tuple[bytes, list[tuple]]:
    """Render a clean resume. Returns (pdf bytes, list of free rectangles usable for hiding text)."""
    s = SAMPLES[sample_id]
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    page.draw_rect(pymupdf.Rect(0, 0, PAGE_W, 96), color=None, fill=(0.93, 0.95, 0.99))
    page.insert_text((MARGIN, 52), s["name"], fontsize=22, fontname="hebo", color=(0.1, 0.12, 0.2))
    page.insert_text((MARGIN, 74), f"{s['title']}  |  {s['contact']}", fontsize=9.5, color=(0.25, 0.27, 0.35))
    y = 128
    for head, lines in s["sections"].items():
        page.insert_text((MARGIN, y), head.upper(), fontsize=11, fontname="hebo", color=(0.15, 0.3, 0.6))
        page.draw_line((MARGIN, y + 4), (PAGE_W - MARGIN, y + 4), color=(0.8, 0.83, 0.9), width=0.6)
        y += 20
        for line in lines:
            rect = pymupdf.Rect(MARGIN, y - 10, PAGE_W - MARGIN, y + 30)
            used = page.insert_textbox(rect, line, fontsize=10, color=(0.12, 0.12, 0.15))
            y += 15 if used >= 14 else 28
        y += 12
    if design_grey:
        page.insert_text((MARGIN, y + 6), "References available on request", fontsize=9, color=(0.55, 0.55, 0.58))
    if small_print:
        page.insert_text((MARGIN, PAGE_H - 24), "This resume was generated from a template. Last updated 2026.",
                         fontsize=6, color=(0.45, 0.45, 0.45))
    free = [
        (MARGIN, y + 20, PAGE_W - MARGIN, PAGE_H - 60),   # space below content
        (MARGIN, PAGE_H - 50, PAGE_W - MARGIN, PAGE_H - 30),  # bottom margin
        (PAGE_W - MARGIN + 4, 120, PAGE_W - 6, 400),      # right margin
        (300, 22, PAGE_W - 20, 56),                        # header band (light blue background)
    ]
    data = doc.tobytes(garbage=3, deflate=True)
    doc.close()
    return data, free


def make_docx(sample_id: str) -> docx.Document:
    s = SAMPLES[sample_id]
    d = docx.Document()
    h = d.add_paragraph()
    r = h.add_run(s["name"])
    r.bold, r.font.size = True, Pt(20)
    d.add_paragraph(f"{s['title']}  |  {s['contact']}")
    for head, lines in s["sections"].items():
        p = d.add_paragraph()
        hr = p.add_run(head.upper())
        hr.bold, hr.font.color.rgb = True, RGBColor(0x26, 0x4D, 0x99)
        for line in lines:
            d.add_paragraph(line)
    return d


def docx_bytes(d: docx.Document) -> bytes:
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
